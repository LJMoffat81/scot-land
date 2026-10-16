"""RoS C14 counts to v1.4 DZ2011 inputs (O-D1).

Reads year, data zone, data zone code and volume only. Price columns are never
requested. Zone status is fixed from sales counts, suppression and lookup flags
before any share list is attached. ``unavailable`` is set only by the
pre-registered cost-availability rule.
"""

from __future__ import annotations

import csv
import hashlib
import json
import logging
from collections import defaultdict
from pathlib import Path
from typing import Mapping, Sequence

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.cell.read_only import EmptyCell

from georgist import surface_rent as sr

logger = logging.getLogger(__name__)

SCOTLAND_CODE = "S92000003"
SCORING_YEAR = 2024
DZ2011_MIN = "S01006506"
DZ2011_MAX = "S01013481"
EXPECTED_LOOKUP_CODES = 7392
# Official DZ2022 lookup census. Runs only when the lookup has all 7,392 codes.
REAL_COLLISION_NAMES = 47
REAL_COLLISION_CODES = 94
REAL_DOUBLE_ROW_CODES = 92
REAL_DOUBLE_ROW_NAMES = 46
MUIRHOUSE_04_NAME = "Muirhouse - 04"
MUIRHOUSE_04_CODES = frozenset({"S01014990", "S01018619"})
REAL_ABSENT_FROM_C14 = 8
REAL_EXCLUDED_DZ2011 = 124
REAL_EXCLUDED_COLLISION = 113
REAL_EXCLUDED_ABSENT = 11
REAL_EXCLUDED_IN_SCORED_SET = 109
REAL_SCORED_SET = 6188
CAUSE_COLLISION = "ros_name_collision"
CAUSE_ABSENT = "absent_from_c14"
CAUSE_ORDINARY = "ordinary_suppression"
NAMED_EXCLUSION_CAUSES = frozenset({CAUSE_COLLISION, CAUSE_ABSENT})
# Statistician's excluded-set list. Hashed before it is parsed.
EXCLUSION_LIST_SHA256 = (
    "13a870fe7ef7118e759d1aa8f46724b4c75c0afb6f7033ad9888bc168a2a78ba"
)
DEFAULT_EXCLUSION_LIST = Path(
    "/cursor/stores/bc-6e71cd70-dc10-4637-afcb-780b94d8b31e/"
    "ros-c14-input-2/collision_excluded_dz2011.csv"
)
COST_AVAILABILITY_RULE = "pre-registered cost-availability rule residual_land_value"

# Header text from C14. "Data zone" is the name column.
BLIND_COLUMNS = (
    "calendar year",
    "data zone",
    "data zone code",
    "volume",
)
# Never requested. Named so a test can assert they stay out of the reader.
PRICE_COLUMNS = (
    "median",
    "lower quartile",
    "upper quartile",
    "total value",
)

PINNED_CROSSWALK_SHA256 = (
    "7690bcbebf9968ca9b02cf898832aa9be8fca56f3eb9af70b20e25b5d23e13e4"
)
PINNED_MERGE_MAP_SHA256 = (
    "581b0aa1647c5af7ce5b9617e35fcfa73982bc8712a42eb63242243ee6a95947"
)
PINNED_FROZEN_PLAN_SHA256_PREFIX = "e1ab6d0a"

E8_MIN_SHARE = 0.01
E8_MIN_UPRN = 5
INCOMPLETE_SHARE = 0.01
INCOMPLETE_ADDRESSES = 5

PRINTOUT_STATUSES = (
    "scored",
    "unavailable",
    "incomplete_suppressed",
    "absent_from_c14",
)


class InputBuildError(RuntimeError):
    """Loud failure: non-zero exit, no output file."""


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def requested_columns(scored_run: bool) -> tuple[str, ...]:
    """Columns the reader asks for. Price columns are never included.

    ``scored_run`` does not widen the request. The parameter is kept so callers
    can show that both modes ask for the same four fields.
    """
    del scored_run
    return BLIND_COLUMNS


def _norm_header(text: str) -> str:
    return " ".join(str(text).strip().lower().split())


def _in_dz2011_range(code: str) -> bool:
    return DZ2011_MIN <= code <= DZ2011_MAX


def _parse_year(raw: str) -> int | None:
    text = str(raw).strip()
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _parse_volume(raw: str) -> tuple[int | None, bool]:
    """Return (volume, suppressed). Markers [c] and [z] are suppressed."""
    text = str(raw).strip()
    if text in {"[c]", "[z]"} or "[c]" in text or "[z]" in text:
        return None, True
    if not text:
        raise InputBuildError("volume cell is empty")
    try:
        volume = int(float(text))
    except ValueError as exc:
        raise InputBuildError("volume cell is not a count or a suppression marker") from exc
    return volume, volume < 5


def _header_index(cells: Sequence[str]) -> dict[str, int] | None:
    normalised = [_norm_header(cell) for cell in cells]
    if "calendar year" not in normalised or "data zone code" not in normalised:
        return None
    if "volume" not in normalised or "data zone" not in normalised:
        return None
    return {name: index for index, name in enumerate(normalised) if name}


def _cell_text(cell: object) -> str:
    """Text of a real cell. Empty and merged fillers contribute nothing."""
    if cell is None or isinstance(cell, (EmptyCell, MergedCell)):
        return ""
    raw = getattr(cell, "value", None)
    return "" if raw is None else str(raw)


# Pair identity uses these blind columns. The zone code is excluded so a
# cartesian copy is one row. Price columns are not inputs.
_KEY_COLUMNS = tuple(name for name in BLIND_COLUMNS if name != "data zone code")


def _blind_hash(projected: Mapping[str, str]) -> str:
    """Hash of the blind columns except the zone code. Nothing else is included."""
    digest = hashlib.sha256()
    for name in _KEY_COLUMNS:
        digest.update(name.encode("utf-8"))
        digest.update(b"\0")
        digest.update(str(projected.get(name, "")).encode("utf-8"))
        digest.update(b"\0")
    return digest.hexdigest()


def _project(fields: Sequence[str], index_by_name: Mapping[str, int], names: Sequence[str]) -> dict[str, str]:
    """Copy only the requested columns. Other cells are not stored."""
    projected: dict[str, str] = {}
    for name in names:
        index = index_by_name[name]
        if index >= len(fields):
            projected[name] = ""
            continue
        value = fields[index]
        projected[name] = "" if value is None else str(value).strip()
    return projected


def _project_xlsx_row(row: Sequence[object], index_by_name: Mapping[str, int], names: Sequence[str]) -> dict[str, str]:
    """Read blind columns only. Other cells in the row are not touched."""
    projected: dict[str, str] = {}
    for name in names:
        index = index_by_name[name]
        if index >= len(row):
            projected[name] = ""
            continue
        projected[name] = _cell_text(row[index]).strip()
    return projected


def _read_csv_rows(path: Path, names: Sequence[str]) -> list[dict[str, str]]:
    index_by_name: dict[str, int] | None = None
    projected: list[dict[str, str]] = []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        for cells in reader:
            found = _header_index(cells)
            if found is not None and all(name in found for name in names):
                index_by_name = found
                break
        if index_by_name is None:
            raise InputBuildError("C14 header row not found")
        for cells in reader:
            item = _project(cells, index_by_name, names)
            if not any(item[name].strip() for name in names):
                continue
            item["row_identity"] = _blind_hash(item)
            projected.append(item)
    return projected


def _read_xlsx_rows(path: Path, names: Sequence[str]) -> list[dict[str, str]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "C14" not in workbook.sheetnames:
            raise InputBuildError("workbook has no C14 sheet")
        sheet = workbook["C14"]
        index_by_name: dict[str, int] | None = None
        projected: list[dict[str, str]] = []
        for row in sheet.iter_rows():
            if index_by_name is None:
                texts = [_cell_text(cell) for cell in row]
                found = _header_index(texts)
                if found is not None and all(name in found for name in names):
                    index_by_name = found
                continue
            item = _project_xlsx_row(row, index_by_name, names)
            if not any(item[name].strip() for name in names):
                continue
            item["row_identity"] = _blind_hash(item)
            projected.append(item)
        if index_by_name is None:
            raise InputBuildError("C14 header row not found")
        return projected
    finally:
        workbook.close()


def read_c14(path: Path, *, scored_run: bool) -> list[dict[str, str]]:
    """Read C14 with column selection. Price columns are not requested."""
    names = requested_columns(scored_run)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv_rows(path, names)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx_rows(path, names)
    raise InputBuildError(f"unsupported C14 file type: {suffix}")


def _load_lookup(path: Path) -> list[dict[str, str]]:
    """Official DZ2022 codes and names. Collisions come from these names."""
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fields = {_norm_header(name): name for name in (reader.fieldnames or [])}
        code_key = (
            fields.get("code")
            or fields.get("dz2022")
            or fields.get("data zone code")
            or fields.get("geography code")
        )
        name_key = (
            fields.get("name")
            or fields.get("data zone")
            or fields.get("zone name")
            or fields.get("dzname")
            or fields.get("datazonename")
        )
        if code_key is None:
            raise InputBuildError("lookup file has no zone code column")
        if name_key is None:
            raise InputBuildError("lookup file has no name column")
        rows: list[dict[str, str]] = []
        for row in reader:
            code = str(row.get(code_key) or "").strip().upper()
            if not code:
                continue
            name = str(row.get(name_key) or "").strip()
            rows.append({"code": code, "name": name})
    codes = [row["code"] for row in rows]
    if len(set(codes)) != len(codes):
        raise InputBuildError("lookup file repeats a zone code")
    return rows


def _collision_sets(
    lookup_rows: Sequence[Mapping[str, str]],
) -> tuple[set[str], dict[str, set[str]], dict[str, str]]:
    """Names held by more than one lookup code. C14 names are not used."""
    codes_for_name: dict[str, set[str]] = defaultdict(set)
    name_for_code: dict[str, str] = {}
    for row in lookup_rows:
        code = str(row["code"])
        name = str(row.get("name") or "").strip()
        name_for_code[code] = name
        if name:
            codes_for_name[name].add(code)
    collision_names = {name for name, codes in codes_for_name.items() if len(codes) > 1}
    pairs = {name: set(codes_for_name[name]) for name in collision_names}
    collision_codes = {code for codes in pairs.values() for code in codes}
    return collision_codes, pairs, name_for_code


def _load_crosswalk(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fields = {_norm_header(name): name for name in (reader.fieldnames or [])}
        col_i = fields.get("dz2011") or fields.get("dz2011_code")
        col_j = fields.get("dz2022") or fields.get("dz2022_code")
        col_n = fields.get("uprn_count") or fields.get("n")
        if not col_i or not col_j or not col_n:
            raise InputBuildError("crosswalk needs dz2011, dz2022 and uprn_count")
        pairs: list[dict] = []
        for row in reader:
            code_i = str(row.get(col_i) or "").strip().upper()
            code_j = str(row.get(col_j) or "").strip().upper()
            if not code_i or not code_j:
                continue
            try:
                count = int(float(str(row.get(col_n) or "").strip()))
            except ValueError as exc:
                raise InputBuildError("crosswalk uprn_count is not an integer") from exc
            if count < 0:
                raise InputBuildError("crosswalk uprn_count is negative")
            pairs.append({"dz2011": code_i, "dz2022": code_j, "uprn_count": count})
    if not pairs:
        raise InputBuildError("crosswalk is empty")
    return pairs


def _load_merge_map(path: Path) -> list[str]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fields = {_norm_header(name): name for name in (reader.fieldnames or [])}
        col = fields.get("dz2011") or fields.get("dz2011_code") or fields.get("code")
        if col is None:
            raise InputBuildError("merge map has no dz2011 column")
        codes = []
        for row in reader:
            code = str(row.get(col) or "").strip().upper()
            if code:
                codes.append(code)
    return sorted(set(codes))


def _load_ur(path: Path) -> dict[str, dict]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fields = {_norm_header(name): name for name in (reader.fieldnames or [])}
        col = fields.get("dz2011") or fields.get("code") or fields.get("data zone code")
        ur_col = fields.get("ur6") or fields.get("ur6_class") or fields.get("class")
        if col is None or ur_col is None:
            raise InputBuildError("urban-rural lookup needs dz2011 and ur6")
        out: dict[str, dict] = {}
        for row in reader:
            code = str(row.get(col) or "").strip().upper()
            if not code:
                continue
            try:
                ur6 = int(float(str(row.get(ur_col) or "").strip()))
            except ValueError as exc:
                raise InputBuildError("urban-rural class is not an integer") from exc
            out[code] = {"ur6": ur6}
    return out


def _e8_keep(pairs: Sequence[Mapping]) -> list[dict]:
    """Drop slivers with j as the source zone, then renormalise v_ij on j."""
    by_j: dict[str, list[Mapping]] = defaultdict(list)
    for pair in pairs:
        by_j[str(pair["dz2022"])].append(pair)
    kept: list[dict] = []
    for edges in by_j.values():
        total = sum(int(edge["uprn_count"]) for edge in edges)
        if total <= 0:
            continue
        survivors = []
        for edge in edges:
            count = int(edge["uprn_count"])
            share = count / total
            if share < E8_MIN_SHARE and count < E8_MIN_UPRN:
                continue
            survivors.append(edge)
        kept_total = sum(int(edge["uprn_count"]) for edge in survivors)
        if kept_total <= 0:
            continue
        for edge in survivors:
            count = int(edge["uprn_count"])
            kept.append(
                {
                    "dz2011": edge["dz2011"],
                    "dz2022": edge["dz2022"],
                    "uprn_count": count,
                    "v": count / kept_total,
                }
            )
    return kept


def _case_labels(kept: Sequence[Mapping]) -> dict[str, str]:
    by_i: dict[str, set[str]] = defaultdict(set)
    by_j: dict[str, set[str]] = defaultdict(set)
    for edge in kept:
        by_i[str(edge["dz2011"])].add(str(edge["dz2022"]))
        by_j[str(edge["dz2022"])].add(str(edge["dz2011"]))
    labels: dict[str, str] = {}
    for code_i, sources in by_i.items():
        if len(sources) == 1:
            source = next(iter(sources))
            labels[code_i] = "1:1" if len(by_j[source]) == 1 else "split"
        elif any(len(by_j[source]) > 1 for source in sources):
            labels[code_i] = "split"
        else:
            labels[code_i] = "merge"
    return labels


def _check_pins(
    *,
    check_pins: bool,
    frozen_plan_path: Path | None,
    crosswalk_path: Path,
    merge_map_path: Path,
    lookup_codes: Sequence[str],
) -> dict[str, str]:
    frozen_sha = file_sha256(frozen_plan_path) if frozen_plan_path and frozen_plan_path.is_file() else ""
    crosswalk_sha = file_sha256(crosswalk_path)
    merge_sha = file_sha256(merge_map_path)
    pins = {
        "frozen_plan_sha256": frozen_sha,
        "crosswalk_sha256": crosswalk_sha,
        "merge_map_sha256": merge_sha,
        "frozen_plan_sha256_prefix_expected": PINNED_FROZEN_PLAN_SHA256_PREFIX,
        "crosswalk_sha256_expected": PINNED_CROSSWALK_SHA256,
        "merge_map_sha256_expected": PINNED_MERGE_MAP_SHA256,
    }
    if not check_pins:
        return pins
    if not frozen_sha.startswith(PINNED_FROZEN_PLAN_SHA256_PREFIX):
        raise InputBuildError("frozen plan sha does not match")
    if crosswalk_sha != PINNED_CROSSWALK_SHA256:
        raise InputBuildError("crosswalk sha does not match")
    if merge_sha != PINNED_MERGE_MAP_SHA256:
        raise InputBuildError("merge map sha does not match")
    if len(lookup_codes) != EXPECTED_LOOKUP_CODES:
        raise InputBuildError("lookup code count does not match")
    return pins


def _vintage_counts(rows: Sequence[Mapping], lookup: set[str]) -> dict[str, int]:
    all_codes: set[str] = set()
    year_codes: set[str] = set()
    for row in rows:
        code = str(row["code"])
        all_codes.add(code)
        if row["year"] == SCORING_YEAR:
            year_codes.add(code)
    return {
        "dz2011_range_codes_all_years": sum(1 for code in all_codes if _in_dz2011_range(code)),
        "dz2011_range_codes_scoring_year": sum(1 for code in year_codes if _in_dz2011_range(code)),
        "codes_not_in_lookup_all_years": sum(1 for code in all_codes if code not in lookup),
        "codes_not_in_lookup_scoring_year": sum(1 for code in year_codes if code not in lookup),
    }


def _assert_vintage(counts: Mapping[str, int]) -> None:
    range_bad = counts["dz2011_range_codes_all_years"] + counts["dz2011_range_codes_scoring_year"]
    lookup_bad = counts["codes_not_in_lookup_all_years"] + counts["codes_not_in_lookup_scoring_year"]
    if range_bad or lookup_bad:
        raise InputBuildError(
            "vintage check failed "
            f"dz2011_range_all_years={counts['dz2011_range_codes_all_years']} "
            f"dz2011_range_scoring_year={counts['dz2011_range_codes_scoring_year']} "
            f"not_in_lookup_all_years={counts['codes_not_in_lookup_all_years']} "
            f"not_in_lookup_scoring_year={counts['codes_not_in_lookup_scoring_year']}"
        )


def _prepare_rows(raw_rows: Sequence[Mapping[str, str]]) -> tuple[list[dict], dict]:
    prepared: list[dict] = []
    scotland_volume = None
    for raw in raw_rows:
        year = _parse_year(raw.get("calendar year", ""))
        if year is None:
            continue
        code = str(raw.get("data zone code", "")).strip().upper()
        name = str(raw.get("data zone", "")).strip()
        volume, suppressed = _parse_volume(raw.get("volume", ""))
        if code == SCOTLAND_CODE:
            if year == SCORING_YEAR:
                if scotland_volume is not None:
                    raise InputBuildError("more than one Scotland total row")
                scotland_volume = 0 if volume is None else volume
            continue
        if not code:
            raise InputBuildError("C14 row has no data zone code")
        prepared.append(
            {
                "year": year,
                "name": name,
                "code": code,
                "volume": volume,
                "suppressed": suppressed,
                "volume_text": str(raw.get("volume", "")).strip(),
                "row_identity": raw.get("row_identity"),
            }
        )
    if scotland_volume is None:
        raise InputBuildError("Scotland total row is missing")
    return prepared, {"scotland_row_volume": scotland_volume}


def _dz2022_status(
    year_rows: Sequence[Mapping],
    lookup_rows: Sequence[Mapping[str, str]],
) -> tuple[dict[str, dict], dict]:
    by_code: dict[str, list[Mapping]] = defaultdict(list)
    for row in year_rows:
        by_code[str(row["code"])].append(row)
    collision_codes, pairs, name_for_code = _collision_sets(lookup_rows)
    lookup_codes = [str(row["code"]) for row in lookup_rows]

    present = set(by_code)
    zones: dict[str, dict] = {}
    for code in lookup_codes:
        if code in collision_codes:
            status = "ros_name_collision"
            suppressed = False
            volume_below_5 = False
        elif code not in present:
            status = "absent_from_c14"
            suppressed = False
            volume_below_5 = False
        else:
            rows = by_code[code]
            volume_below_5 = any(bool(row["suppressed"]) for row in rows)
            marker = any(row["volume"] is None for row in rows)
            if volume_below_5 or marker:
                status = "suppressed"
                suppressed = True
            else:
                status = "ok"
                suppressed = False
        volume_sum = 0
        for row in by_code.get(code, []):
            if row["volume"] is not None and not row["suppressed"] and status == "ok":
                volume_sum += int(row["volume"])
        zones[code] = {
            "code": code,
            "status": status,
            "volume_below_5": volume_below_5 and status == "suppressed",
            "suppressed": suppressed,
            "count": volume_sum if status == "ok" else 0,
            "names": [name_for_code[code]] if name_for_code.get(code) else [],
            "c14_rows": len(by_code.get(code, [])),
        }

    # One sum per lookup pair. The key is a hash of the blind columns except
    # the zone code, so a cartesian copy is not doubled and price cells are unused.
    seen: set[tuple] = set()
    pair_volume: dict[str, int] = {name: 0 for name in pairs}
    raw_sum = 0
    for row in year_rows:
        if row["volume"] is not None:
            raw_sum += int(row["volume"])
        code = str(row["code"])
        if code not in collision_codes:
            continue
        pair_name = name_for_code[code]
        identity = row.get("row_identity")
        if not isinstance(identity, str) or len(identity) != 64:
            raise InputBuildError("C14 row has no content key")
        key = (pair_name, identity)
        if key in seen:
            continue
        seen.add(key)
        if row["volume"] is not None:
            pair_volume[pair_name] += int(row["volume"])
    pair_sum = sum(pair_volume.values())
    extra_codes = sorted(set(by_code) - set(lookup_codes))
    national = {
        "raw_dz_sum": raw_sum,
        "pair_dedup_sum": pair_sum,
        "pair_volume": pair_volume,
        "collision_names": len(pairs),
        "collision_codes": len(collision_codes),
        "collision_pairs": {name: sorted(codes) for name, codes in pairs.items()},
        "rows_by_code": {code: len(rows) for code, rows in by_code.items()},
        "extra_codes": extra_codes,
    }
    return zones, national


def _addresses_trip(count: int, total: int) -> bool:
    """True when this address group alone is above 1% or at least 5 addresses."""
    if count <= 0 or total <= 0:
        return False
    return (count / total) > INCOMPLETE_SHARE or count >= INCOMPLETE_ADDRESSES


def _exclusion_cause(collision_n: int, absent_n: int, total_n: int, incomplete: bool) -> str:
    """Collision, then absent codes, then ordinary suppression (fewer than 5 sales)."""
    if _addresses_trip(collision_n, total_n):
        return CAUSE_COLLISION
    if _addresses_trip(absent_n, total_n):
        return CAUSE_ABSENT
    if incomplete:
        return CAUSE_ORDINARY
    return ""


def load_pinned_exclusion_list(path: Path) -> dict[str, dict]:
    """Hash the exclusion list before reading any row. Mismatch raises."""
    if not path.is_file():
        raise InputBuildError(f"exclusion list missing: {path}")
    digest = file_sha256(path)
    if digest != EXCLUSION_LIST_SHA256:
        raise InputBuildError(f"exclusion list sha256 {digest} != {EXCLUSION_LIST_SHA256}")
    pinned: dict[str, dict] = {}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fields = {_norm_header(name): name for name in (reader.fieldnames or [])}
        code_col = fields.get("dz2011_code") or fields.get("dz2011")
        cause_col = fields.get("cause")
        scored_col = fields.get("in_scored_set")
        if not code_col or not cause_col or not scored_col:
            raise InputBuildError("exclusion list is missing dz2011_code, cause or in_scored_set")
        for row in reader:
            code = str(row.get(code_col) or "").strip().upper()
            if not code:
                continue
            cause = str(row.get(cause_col) or "").strip()
            scored_text = str(row.get(scored_col) or "").strip().lower()
            if scored_text not in {"true", "false"}:
                raise InputBuildError(f"{code} in_scored_set is not true or false")
            pinned[code] = {"cause": cause, "in_scored_set": scored_text == "true"}
    return pinned


def diff_exclusion_sets(
    observed: Mapping[str, Mapping],
    pinned: Mapping[str, Mapping],
    *,
    ordinary_codes: set[str] | None = None,
) -> list[str]:
    """Compare codes, cause and scored-set membership. Ordinary suppression must not overlap."""
    errors: list[str] = []
    ordinary = ordinary_codes or set()
    overlap = (set(observed) | set(pinned)) & ordinary
    if overlap:
        errors.append(
            "ordinary suppression overlaps excluded set: " + ",".join(sorted(overlap))
        )
    missing = sorted(set(pinned) - set(observed))
    extra = sorted(set(observed) - set(pinned))
    if missing:
        errors.append("exclusion list missing codes: " + ",".join(missing))
    if extra:
        errors.append("exclusion list extra codes: " + ",".join(extra))
    for code in sorted(set(observed) & set(pinned)):
        if observed[code]["cause"] != pinned[code]["cause"]:
            errors.append(
                f"{code} cause {observed[code]['cause']} != {pinned[code]['cause']}"
            )
        if bool(observed[code]["in_scored_set"]) != bool(pinned[code]["in_scored_set"]):
            errors.append(f"{code} in_scored_set mismatch")
    return errors


def exclusion_list_errors(zones: Sequence[Mapping], path: Path) -> list[str]:
    """Diff collision and absent exclusions. Ordinary suppression stays out of the set."""
    try:
        pinned = load_pinned_exclusion_list(path)
    except InputBuildError as exc:
        return [str(exc)]
    observed = {
        str(zone["dz2011"]): {
            "cause": zone["exclusion_cause"],
            "in_scored_set": bool(zone["in_scored_set"]),
        }
        for zone in zones
        if zone.get("exclusion_cause") in NAMED_EXCLUSION_CAUSES
    }
    ordinary = {
        str(zone["dz2011"])
        for zone in zones
        if zone.get("exclusion_cause") == CAUSE_ORDINARY
    }
    return diff_exclusion_sets(observed, pinned, ordinary_codes=ordinary)


def real_census_errors(observed: Mapping, *, lookup_count: int) -> list[str]:
    """Compare a full DZ2022 run with the fixed census. Smaller lookups skip."""
    if lookup_count != EXPECTED_LOOKUP_CODES:
        return []
    expected = {
        "collision_names": REAL_COLLISION_NAMES,
        "collision_codes": REAL_COLLISION_CODES,
        "double_row_codes": REAL_DOUBLE_ROW_CODES,
        "double_row_names": REAL_DOUBLE_ROW_NAMES,
        "absent_from_c14": REAL_ABSENT_FROM_C14,
        "excluded_dz2011": REAL_EXCLUDED_DZ2011,
        "excluded_collision": REAL_EXCLUDED_COLLISION,
        "excluded_absent_from_c14": REAL_EXCLUDED_ABSENT,
        "excluded_in_scored_set": REAL_EXCLUDED_IN_SCORED_SET,
        "scored_set": REAL_SCORED_SET,
    }
    errors = [
        f"{key} {observed.get(key)} != {value}"
        for key, value in expected.items()
        if observed.get(key) != value
    ]
    muirhouse_codes = set(observed.get("muirhouse_04_codes") or [])
    if muirhouse_codes != set(MUIRHOUSE_04_CODES):
        errors.append("Muirhouse - 04 codes are not S01014990 and S01018619")
    if observed.get("muirhouse_04_name") != MUIRHOUSE_04_NAME:
        errors.append("Muirhouse - 04 name does not match the lookup")
    if observed.get("muirhouse_04_rows") != 1:
        errors.append("Muirhouse - 04 does not have one row attached to both codes")
    collision = observed.get("excluded_collision")
    absent = observed.get("excluded_absent_from_c14")
    if (
        isinstance(collision, int)
        and isinstance(absent, int)
        and collision + absent != observed.get("excluded_dz2011")
    ):
        errors.append(
            "excluded_dz2011 is not ros_name_collision plus absent_from_c14 "
            "(ordinary suppression is separate)"
        )
    return errors


def _observed_census(
    *,
    zones_2022: Mapping[str, Mapping],
    national_counts: Mapping,
    dz2011: Sequence[Mapping],
    scored_set: set[str],
) -> dict:
    pairs = national_counts["collision_pairs"]
    rows_by_code = national_counts["rows_by_code"]
    double_codes = [
        code
        for name, codes in pairs.items()
        for code in codes
        if rows_by_code.get(code, 0) == 2
    ]
    double_names = {
        name
        for name, codes in pairs.items()
        if codes and all(rows_by_code.get(code, 0) == 2 for code in codes)
    }
    muirhouse_codes = pairs.get(MUIRHOUSE_04_NAME, set())
    muirhouse_rows = {rows_by_code.get(code, 0) for code in MUIRHOUSE_04_CODES}
    collision_excluded = [
        zone for zone in dz2011 if zone.get("exclusion_cause") == CAUSE_COLLISION
    ]
    absent_excluded = [zone for zone in dz2011 if zone.get("exclusion_cause") == CAUSE_ABSENT]
    named = collision_excluded + absent_excluded
    ordinary = [zone for zone in dz2011 if zone.get("exclusion_cause") == CAUSE_ORDINARY]
    return {
        "collision_names": national_counts["collision_names"],
        "collision_codes": national_counts["collision_codes"],
        "double_row_codes": len(double_codes),
        "double_row_names": len(double_names),
        "muirhouse_04_codes": set(muirhouse_codes),
        "muirhouse_04_name": MUIRHOUSE_04_NAME if muirhouse_codes else "",
        "muirhouse_04_rows": next(iter(muirhouse_rows)) if len(muirhouse_rows) == 1 else None,
        "absent_from_c14": sum(
            1 for zone in zones_2022.values() if zone["status"] == "absent_from_c14"
        ),
        "excluded_dz2011": len(named),
        "excluded_collision": len(collision_excluded),
        "excluded_absent_from_c14": len(absent_excluded),
        "excluded_ordinary_suppression": len(ordinary),
        "excluded_in_scored_set": sum(1 for zone in named if zone["in_scored_set"]),
        "scored_set": len(scored_set),
    }


def _dz2011_rows(
    *,
    lookup_zones: Mapping[str, Mapping],
    pairs: Sequence[Mapping],
    scored_set: set[str],
    ur: Mapping[str, Mapping],
) -> list[dict]:
    raw_by_i: dict[str, list[Mapping]] = defaultdict(list)
    for pair in pairs:
        raw_by_i[str(pair["dz2011"])].append(pair)
    kept = _e8_keep(pairs)
    cases = _case_labels(kept)
    kept_by_i: dict[str, list[Mapping]] = defaultdict(list)
    for edge in kept:
        kept_by_i[str(edge["dz2011"])].append(edge)

    codes = set(raw_by_i) | set(scored_set) | set(ur)
    rows: list[dict] = []
    for code in sorted(codes):
        raw_edges = raw_by_i.get(code, [])
        total_n = sum(int(edge["uprn_count"]) for edge in raw_edges)
        absent_n = 0
        low_n = 0
        collision_n = 0
        affected_n = 0
        for edge in raw_edges:
            count = int(edge["uprn_count"])
            source = lookup_zones.get(str(edge["dz2022"]))
            if source is None or source["status"] != "ok":
                affected_n += count
            if source is None or source["status"] == "absent_from_c14":
                absent_n += count
            elif source["status"] == "suppressed":
                low_n += count
            elif source["status"] == "ros_name_collision":
                collision_n += count
        share = (affected_n / total_n) if total_n else 0.0
        incomplete = total_n > 0 and (share > INCOMPLETE_SHARE or affected_n >= INCOMPLETE_ADDRESSES)
        pure_absent = absent_n > 0 and low_n == 0 and collision_n == 0 and affected_n == absent_n
        exclusion_cause = _exclusion_cause(collision_n, absent_n, total_n, incomplete)
        ur_row = ur.get(code) or {}
        if incomplete:
            status = "incomplete_suppressed"
        elif pure_absent:
            status = "absent_from_c14"
        elif code in scored_set:
            status = "scored"
        else:
            status = "not_in_scored_set"
        count_s = 0.0
        if status == "scored":
            for edge in kept_by_i.get(code, []):
                source = lookup_zones.get(str(edge["dz2022"]))
                if source is None or source["status"] != "ok":
                    continue
                count_s += float(edge["v"]) * float(source["count"])
        in_scored_set = code in scored_set
        passed = in_scored_set and status in {"scored", "unavailable"}
        absent_flag = pure_absent or status == "absent_from_c14"
        rows.append(
            {
                "dz2011": code,
                "status": status,
                "absent_from_c14": absent_flag,
                "volume_below_5": low_n > 0 and absent_n == 0 and collision_n == 0,
                "incomplete_suppressed": status == "incomplete_suppressed",
                "exclusion_cause": exclusion_cause,
                "unavailable": status == "unavailable",
                "case": cases.get(code, "not_applicable"),
                "count_s": count_s,
                "in_scored_set": in_scored_set,
                "passed_to_scorer": passed,
                "ur6": ur_row.get("ur6"),
                "medians_quartiles": "not_crosswalkable",
                "affected_addresses": affected_n,
                "address_total": total_n,
            }
        )
    return rows


def _counts_by_status(zones: Sequence[Mapping]) -> dict[str, int]:
    counts = {name: 0 for name in PRINTOUT_STATUSES}
    for zone in zones:
        status = str(zone["status"])
        if status in counts:
            counts[status] += 1
        if zone.get("absent_from_c14") and status != "absent_from_c14":
            counts["absent_from_c14"] += 1
    return counts


def _counts_by_ur(zones: Sequence[Mapping]) -> dict[str, dict[str, int]]:
    table = {name: {str(klass): 0 for klass in range(1, 7)} for name in PRINTOUT_STATUSES}
    for zone in zones:
        ur6 = zone.get("ur6")
        if ur6 not in range(1, 7):
            continue
        key = str(int(ur6))
        status = str(zone["status"])
        if status in table:
            table[status][key] += 1
        if zone.get("absent_from_c14") and status != "absent_from_c14":
            table["absent_from_c14"][key] += 1
    return table


def render_printout(result: Mapping) -> str:
    """Counts and flags only."""
    lines = [
        "dz2022 "
        + " ".join(f"{key}={result['dz2022_counts'][key]}" for key in sorted(result["dz2022_counts"])),
        "dz2011 "
        + " ".join(f"{key}={result['dz2011_counts'][key]}" for key in PRINTOUT_STATUSES),
        "vintage "
        + " ".join(f"{key}={result['vintage_counts'][key]}" for key in sorted(result["vintage_counts"])),
        "lookup_absent_from_c14={0}".format(result["lookup_codes_absent_from_c14"]),
        "national scotland_row_volume={scotland_row_volume} raw_dz_sum={raw_dz_sum} "
        "pair_dedup_sum={pair_dedup_sum}".format(**result["national"]),
        "weights=all_uprn",
    ]
    excluded = result["excluded_incomplete_by_ur6"]
    lines.append(
        "excluded_incomplete_suppressed_by_ur6 "
        + " ".join(f"class{klass}={excluded[str(klass)]}" for klass in range(1, 7))
    )
    for status in PRINTOUT_STATUSES:
        by_class = result["dz2011_by_ur6"][status]
        parts = " ".join(f"class{klass}={by_class[str(klass)]}" for klass in range(1, 7))
        lines.append(f"ur6 {status} {parts}")
    return "\n".join(lines) + "\n"


def _share_list(zones: Sequence[Mapping], estimates: Mapping[str, float] | None) -> list[float] | None:
    """Attach shares after status exists. Unavailable zones enter as 0.0."""
    if estimates is None:
        return None
    shares: list[float] = []
    for zone in zones:
        if not zone["passed_to_scorer"]:
            continue
        if zone["status"] == "unavailable":
            shares.append(0.0)
            continue
        code = str(zone["dz2011"])
        if code not in estimates:
            return None
        shares.append(float(estimates[code]))
    return shares


def load_cost_availability(path: Path) -> dict[str, dict[str, float]]:
    """Price and rebuild cost for the pre-registered rule. Estimate keys are ignored."""
    raw = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(raw, dict):
        raise InputBuildError("cost-availability file must be a JSON object")
    out: dict[str, dict[str, float]] = {}
    for code, row in raw.items():
        if not isinstance(row, dict) or "price" not in row or "drc" not in row:
            raise InputBuildError(f"{code} cost-availability row needs price and drc")
        out[str(code).strip().upper()] = {
            "price": float(row["price"]),
            "drc": float(row["drc"]),
        }
    return out


def _apply_cost_availability(zones: list[dict], costs: Mapping[str, Mapping] | None) -> None:
    """Flip scored zones only when residual_land_value labels them unavailable."""
    if not costs:
        return
    for zone in zones:
        if zone["status"] != "scored":
            continue
        code = str(zone["dz2011"])
        row = costs.get(code)
        if row is None:
            continue
        residual = sr.residual_land_value(float(row["price"]), float(row["drc"]))
        if sr.estimate_label_for_residual(residual) != sr.ESTIMATE_LABEL_UNAVAILABLE:
            continue
        zone["status"] = "unavailable"
        zone["unavailable"] = True
        zone["unavailable_from"] = COST_AVAILABILITY_RULE
        zone["passed_to_scorer"] = bool(zone["in_scored_set"])
        logger.info("%s unavailable_from=%s", code, COST_AVAILABILITY_RULE)


def build_inputs(
    *,
    c14_path: Path,
    lookup_path: Path,
    crosswalk_path: Path,
    merge_map_path: Path,
    ur_path: Path,
    frozen_plan_path: Path | None = None,
    scored_run: bool = False,
    estimates: Mapping[str, float] | None = None,
    cost_availability: Mapping[str, Mapping] | None = None,
    exclusion_list_path: Path | None = None,
    main_commit_sha: str = "",
    check_pins: bool = False,
) -> dict:
    """Build the DZ2011 input document. Status does not read ``estimates``."""
    columns = requested_columns(scored_run)
    raw = read_c14(c14_path, scored_run=scored_run)
    prepared, scotland = _prepare_rows(raw)
    lookup_rows = _load_lookup(lookup_path)
    lookup_codes = [row["code"] for row in lookup_rows]
    lookup_set = set(lookup_codes)
    vintage = _vintage_counts(prepared, lookup_set)
    _assert_vintage(vintage)
    pins = _check_pins(
        check_pins=check_pins,
        frozen_plan_path=frozen_plan_path,
        crosswalk_path=crosswalk_path,
        merge_map_path=merge_map_path,
        lookup_codes=lookup_codes,
    )
    year_rows = [row for row in prepared if row["year"] == SCORING_YEAR]
    zones_2022, national_counts = _dz2022_status(year_rows, lookup_rows)
    national = {
        "scotland_row_volume": scotland["scotland_row_volume"],
        "raw_dz_sum": national_counts["raw_dz_sum"],
        "pair_dedup_sum": national_counts["pair_dedup_sum"],
    }
    pairs = _load_crosswalk(crosswalk_path)
    scored_set = set(_load_merge_map(merge_map_path))
    ur = _load_ur(ur_path)
    dz2011 = _dz2011_rows(
        lookup_zones=zones_2022,
        pairs=pairs,
        scored_set=scored_set,
        ur=ur,
    )
    _apply_cost_availability(dz2011, cost_availability)
    # Estimates are applied only after every status field is set.
    share_list = _share_list(dz2011, estimates) if scored_run else None
    dz2022_counts = {
        "ok": sum(1 for zone in zones_2022.values() if zone["status"] == "ok"),
        "suppressed": sum(1 for zone in zones_2022.values() if zone["status"] == "suppressed"),
        "ros_name_collision": sum(
            1 for zone in zones_2022.values() if zone["status"] == "ros_name_collision"
        ),
        "absent_from_c14": sum(
            1 for zone in zones_2022.values() if zone["status"] == "absent_from_c14"
        ),
    }
    excluded = [
        zone
        for zone in dz2011
        if zone["incomplete_suppressed"] and zone["in_scored_set"]
    ]
    excluded_by_ur = {str(klass): 0 for klass in range(1, 7)}
    for zone in excluded:
        if zone["ur6"] in range(1, 7):
            excluded_by_ur[str(int(zone["ur6"]))] += 1
    observed = _observed_census(
        zones_2022=zones_2022,
        national_counts=national_counts,
        dz2011=dz2011,
        scored_set=scored_set,
    )
    census_errors = real_census_errors(observed, lookup_count=len(lookup_rows))
    if len(lookup_rows) == EXPECTED_LOOKUP_CODES:
        list_path = exclusion_list_path or DEFAULT_EXCLUSION_LIST
        census_errors.extend(exclusion_list_errors(dz2011, list_path))
    if census_errors:
        raise InputBuildError("real-data collision census failed: " + "; ".join(census_errors))
    absent_codes = [
        code for code, zone in zones_2022.items() if zone["status"] == "absent_from_c14"
    ]
    result = {
        "weights": "all_uprn",
        "scoring_year": SCORING_YEAR,
        "vintage": "DZ2022",
        "gbp_display": False,
        "price_columns_read": False,
        "reader_requested_columns": list(columns),
        "disclosures": [
            "Vintage is decided from codes. The C14 header says Data zone.",
            "Collision and absent codes use the existing suppressed treatment. No new fill is added.",
            "The frozen plan does not say whether dates are registration or transaction dates.",
        ],
        "pins": {
            "c14_sha256": file_sha256(c14_path),
            "lookup_sha256": file_sha256(lookup_path),
            "crosswalk_sha256": pins["crosswalk_sha256"],
            "merge_map_sha256": pins["merge_map_sha256"],
            "frozen_plan_sha256": pins["frozen_plan_sha256"],
            "main_commit_sha": main_commit_sha,
            "crosswalk_sha256_expected": pins["crosswalk_sha256_expected"],
            "merge_map_sha256_expected": pins["merge_map_sha256_expected"],
            "frozen_plan_sha256_prefix_expected": pins["frozen_plan_sha256_prefix_expected"],
        },
        "vintage_counts": vintage,
        "lookup_codes_absent_from_c14": len(absent_codes),
        "incomplete_suppressed_dz2011": sum(
            1 for zone in dz2011 if zone["incomplete_suppressed"]
        ),
        "excluded_collision": observed["excluded_collision"],
        "excluded_absent_from_c14": observed["excluded_absent_from_c14"],
        "excluded_ordinary_suppression": observed["excluded_ordinary_suppression"],
        "national": {
            "scotland_row_volume": national["scotland_row_volume"],
            "raw_dz_sum": national["raw_dz_sum"],
            "pair_dedup_sum": national["pair_dedup_sum"],
            "pair_volume": national_counts["pair_volume"],
        },
        "dz2022_counts": dz2022_counts,
        "dz2022": [zones_2022[code] for code in sorted(zones_2022)],
        "dz2011_counts": _counts_by_status(dz2011),
        "dz2011_by_ur6": _counts_by_ur(dz2011),
        "dz2011": dz2011,
        "excluded_incomplete_suppressed": [zone["dz2011"] for zone in excluded],
        "excluded_incomplete_by_ur6": excluded_by_ur,
        "share_list": share_list,
        "collision_names": national_counts["collision_names"],
        "collision_codes": national_counts["collision_codes"],
    }
    result["printout"] = render_printout(result)
    return result


def write_outputs(result: Mapping, out_path: Path, *, shares_path: Path | None = None) -> None:
    """Write the gitignored JSON. Share list only after a scored run."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {key: value for key, value in result.items() if key != "printout"}
    out_path.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    if shares_path is not None and result.get("share_list") is not None:
        shares_path.parent.mkdir(parents=True, exist_ok=True)
        shares_path.write_text(
            json.dumps(result["share_list"]) + "\n",
            encoding="utf-8",
        )
