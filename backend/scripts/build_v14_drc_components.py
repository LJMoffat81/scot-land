#!/usr/bin/env python3
"""Write one DRC component record per DZ2011 zone.

``epc_row`` comes from the EPC rollup, ``nrs2017_row`` from NRS dwellings by
type and rooms, and ``estate_row`` from plot-band shares after
``borrow_plot_bands`` (council x UR6, then UR6 national). No price column
is read. Pins are full sha256 digests.

Usage (from ``backend/``)::

    python -m scripts.build_v14_drc_components --out components.json \\
        --epc epc.csv --nrs2017 nrs.csv --plot-bands plots.csv --ur6 ur6.csv \\
        --n25i dw11.xlsx
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sys
from pathlib import Path

import openpyxl

from georgist import surface_harvester as hv
from georgist import surface_rent as sr
from georgist.v14_real_headers import REAL_HEADERS, header_refusal, read_header_fields

PIN_TABLE = {
    "epc": "13a5f647a833d89ec7aad0a82cd3cfc52d375b31b05e3b5ad7f4f12c998e2618",
    "nrs2017": "b9bd3364ade0a843a3157fe7e33de5003144f18c54f87c9e410c4dd3fc04a9da",
    "plot_bands": "9bb56d3be6181740ae07fb06607c45374ce448d6c73a00cdea3c7c81b4886e48",
    "ur6": "1f32a81532912334e3cd071e2f12f523e66068d14c73e5db4ae51add239b1a92",
    "n25i": "6fdf2951223dd01e5790069d61428ce25d92aed5f260197ee546223113df4e49",
}
PIN_ORDER = ("epc", "nrs2017", "plot_bands", "ur6", "n25i")
FULL_SHA = re.compile(r"^[0-9a-f]{64}$")
COUNCIL_CODE = re.compile(r"^S12\d{6}$")
N25I_SHEET = "2025"
N25I_CODE = "Data Zone code"
N25I_COUNCIL = "Council area code"
REQUIRED_COLUMNS = {
    "epc": ("code",),
    "nrs2017": ("code",),
    "plot_bands": ("code",),
    "ur6": ("code", "ur6_class"),
}


class ComponentsError(ValueError):
    """A DRC component input failed a pin, file, or cell check."""


def copy_pin_table() -> dict[str, str]:
    """Test-only hook. A copy, so production pins stay put."""
    return {key: str(PIN_TABLE[key]) for key in PIN_ORDER}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(path.read_bytes())
    return digest.hexdigest()


def read_price_cell(value: str) -> str:
    """Price-column hook. The component readers do not call it."""
    return value


def borrow_plot_bands(rows, zone_meta, estate_cfg):
    """Council x UR6, then UR6 national. Tests spy on this wrapper."""
    return sr.borrow_plot_bands(rows, zone_meta, estate_cfg)


def _pin_text(value: object) -> str:
    if value is None:
        return "blank"
    text = str(value)
    if text.strip() == "":
        return "blank"
    return text


def assert_pins(paths: dict[str, Path], pins: dict[str, str]) -> None:
    """Full sha256, then the file, then the digest. Readers have not run."""
    for name in PIN_ORDER:
        text = _pin_text(pins.get(name))
        if text == "blank" or FULL_SHA.fullmatch(text) is None:
            raise ComponentsError(
                f"B1_COMPONENTS: REFUSED: pin: expected full sha256, got {text}"
            )
    for name in PIN_ORDER:
        path = paths[name]
        if not path.is_file():
            raise ComponentsError(f"B1_COMPONENTS: ERROR: file: expected {path}, got absent")
        actual = sha256_file(path)
        if actual != pins[name]:
            raise ComponentsError(
                f"B1_COMPONENTS: REFUSED: {name} sha256: expected {pins[name]}, got {actual}"
            )


def _refuse_header(kind: str, header: list[str]) -> None:
    refusal = header_refusal("B1_COMPONENTS", kind, header)
    if refusal is not None:
        raise ComponentsError(refusal)


def parse_component_rows(path: Path, kind: str) -> list[dict[str, str]]:
    """Parse data rows. The header guard has already accepted the header line."""
    text = path.read_text(encoding="utf-8")
    lines = [line for line in text.splitlines() if line.strip() and not line.startswith("#")]
    if not lines:
        raise ComponentsError("B1_COMPONENTS: ERROR: header: expected header, got absent")
    reader = csv.reader(io.StringIO("\n".join(lines) + "\n"))
    parsed = list(reader)
    header = parsed[0]
    width = len(header)
    required = REQUIRED_COLUMNS[kind]
    rows: list[dict[str, str]] = []
    for row_number, cells in enumerate(parsed[1:], start=1):
        if len(cells) < width:
            raise ComponentsError(
                f"B1_COMPONENTS: ERROR: row {row_number}: expected {width} fields, got {len(cells)}"
            )
        item: dict[str, str] = {}
        for index, column in enumerate(header):
            if column == "price":
                continue
            value = cells[index]
            if column in required and value.strip() == "":
                raise ComponentsError(
                    f"B1_COMPONENTS: ERROR: {column} row {row_number}: expected value, got blank"
                )
            item[column] = value
        rows.append(item)
    return rows


def _table(path: Path, kind: str) -> list[dict[str, str]]:
    header = read_header_fields(path)
    if header is None:
        raise ComponentsError("B1_COMPONENTS: ERROR: header: expected header, got absent")
    _refuse_header(kind, header)
    return parse_component_rows(path, kind)


def load_epc(path: Path) -> dict[str, dict[str, str]]:
    return {row["code"]: row for row in _table(path, "epc")}


def load_nrs2017(path: Path) -> dict[str, dict[str, str]]:
    return {row["code"]: row for row in _table(path, "nrs2017")}


def load_plot_bands(path: Path) -> dict[str, dict[str, str]]:
    return {row["code"]: row for row in _table(path, "plot_bands")}


def load_ur6(path: Path) -> dict[str, int]:
    """UR6 class by zone. Council codes come from n25i, not this file."""
    classes: dict[str, int] = {}
    for row_number, row in enumerate(_table(path, "ur6"), start=1):
        code = row["code"]
        if code in classes:
            raise ComponentsError(
                f"B1_COMPONENTS: REFUSED: ur6 code: expected unique, got duplicate {code}"
            )
        text = row["ur6_class"].strip()
        if not text.isdigit() or not 1 <= int(text) <= 6:
            raise ComponentsError(
                f"B1_COMPONENTS: ERROR: ur6_class row {row_number}: expected 1-6, got {text}"
            )
        classes[code] = int(text)
    return classes


def _n25i_sheet(path: Path):
    workbook = openpyxl.load_workbook(path, data_only=True)
    if N25I_SHEET not in workbook.sheetnames:
        workbook.close()
        raise ComponentsError(
            f"B1_COMPONENTS: ERROR: file: expected {path} sheet {N25I_SHEET}, got absent"
        )
    return workbook, workbook[N25I_SHEET]


def _n25i_header_index(path: Path) -> int:
    try:
        return hv._find_header(path, N25I_SHEET)
    except (ValueError, KeyError, OSError):
        return 4


def _stored_width(sheet, row_number: int) -> int:
    columns = [column for (row, column) in sheet._cells if row == row_number]
    if not columns:
        return 0
    return max(columns)


def _row_values(sheet, row_number: int, width: int) -> list[str]:
    values: list[str] = []
    for column in range(1, width + 1):
        cell = sheet.cell(row=row_number, column=column).value
        if cell is None:
            values.append("")
        else:
            values.append(str(cell))
    return values


def _n25i_header_fields(path: Path) -> list[str]:
    header_index = _n25i_header_index(path)
    workbook, sheet = _n25i_sheet(path)
    try:
        return _row_values(sheet, header_index + 1, _stored_width(sheet, header_index + 1))
    finally:
        workbook.close()


def parse_n25i_rows(path: Path) -> list[tuple[str, str]]:
    """Data rows of the n25i sheet. Header acceptance happens in ``load_n25i``."""
    expected = REAL_HEADERS["n25i"]
    width = len(expected)
    _workbook, sheet = _n25i_sheet(path)
    try:
        header_row = _n25i_header_index(path) + 1
        pairs: list[tuple[str, str]] = []
        seen: set[str] = set()
        data_number = 0
        for row_number in range(header_row + 1, sheet.max_row + 1):
            stored = _stored_width(sheet, row_number)
            if stored == 0:
                continue
            data_number += 1
            if stored < width:
                raise ComponentsError(
                    f"B1_COMPONENTS: ERROR: row {data_number}: expected {width} fields, got {stored}"
                )
            values = _row_values(sheet, row_number, width)
            item = {column: values[index] for index, column in enumerate(expected)}
            code = item[N25I_CODE].strip()
            council = item[N25I_COUNCIL].strip()
            if code == "":
                raise ComponentsError(
                    f"B1_COMPONENTS: ERROR: {N25I_CODE} row {data_number}: expected value, got blank"
                )
            if council == "":
                raise ComponentsError(
                    f"B1_COMPONENTS: ERROR: {N25I_COUNCIL} row {data_number}: expected value, got blank"
                )
            if code in seen:
                raise ComponentsError(
                    f"B1_COMPONENTS: REFUSED: n25i code: expected unique, got duplicate {code}"
                )
            seen.add(code)
            if COUNCIL_CODE.fullmatch(council) is None:
                raise ComponentsError(
                    f"B1_COMPONENTS: ERROR: council_code row {data_number}: expected S12 code, got {council}"
                )
            pairs.append((code, council))
        return pairs
    finally:
        _workbook.close()


def load_n25i(path: Path) -> dict[str, str]:
    """Council code by data-zone code. The NRS loader is called unchanged."""
    _refuse_header("n25i", _n25i_header_fields(path))
    pairs = parse_n25i_rows(path)
    frame = hv.load_nrs_dwellings(path, year=2025, vintage="DZ2011")
    councils = {code: council for code, council in pairs}
    if "code" in frame.columns and "council_code" in frame.columns:
        for code, council in zip(frame["code"], frame["council_code"]):
            text = "" if council is None else str(council).strip()
            if text.lower() == "nan" or text == "":
                continue
            councils[str(code)] = text
    return councils


def _zone_meta(classes: dict[str, int], councils: dict[str, str], n25i_path: Path) -> dict[str, tuple[str, int]]:
    meta: dict[str, tuple[str, int]] = {}
    for code, ur6_class in classes.items():
        if code not in councils:
            raise ComponentsError(
                f"B1_COMPONENTS: REFUSED: council_code: expected {code} in n25i {n25i_path}, got absent"
            )
        meta[code] = (councils[code], ur6_class)
    return meta


def build_components(
    paths: dict[str, Path],
    *,
    pins: dict[str, str] | None = None,
) -> list[dict]:
    """One record per UR6 zone. Plot-band shares are borrowed before return."""
    table = PIN_TABLE if pins is None else pins
    assert_pins(paths, table)
    for kind in ("epc", "nrs2017", "plot_bands", "ur6"):
        header = read_header_fields(paths[kind])
        if header is None:
            raise ComponentsError("B1_COMPONENTS: ERROR: header: expected header, got absent")
        _refuse_header(kind, header)
    _refuse_header("n25i", _n25i_header_fields(paths["n25i"]))
    epc = load_epc(paths["epc"])
    nrs = load_nrs2017(paths["nrs2017"])
    plots = load_plot_bands(paths["plot_bands"])
    classes = load_ur6(paths["ur6"])
    councils = load_n25i(paths["n25i"])
    zones = _zone_meta(classes, councils, paths["n25i"])
    estate_cfg = sr.load_config()["estate"]
    estate = borrow_plot_bands(plots, zones, estate_cfg)
    records = []
    for code in sorted(zones):
        records.append(
            {
                "dz2011": code,
                "council_code": zones[code][0],
                "epc_row": epc.get(code),
                "nrs2017_row": nrs.get(code),
                "estate_row": estate.get(code),
            }
        )
    return records


def write_components(path: Path, records: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(records, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--epc", type=Path, required=True)
    parser.add_argument("--nrs2017", type=Path, required=True)
    parser.add_argument("--plot-bands", type=Path, required=True)
    parser.add_argument("--ur6", type=Path, required=True)
    parser.add_argument("--n25i", type=Path, required=True)
    args = parser.parse_args(argv)
    paths = {
        "epc": args.epc,
        "nrs2017": args.nrs2017,
        "plot_bands": args.plot_bands,
        "ur6": args.ur6,
        "n25i": args.n25i,
    }
    try:
        records = build_components(paths)
        write_components(args.out, records)
    except ComponentsError as exc:
        sys.stderr.write(str(exc) + "\n")
        return 1
    sys.stdout.write(f"zones={len(records)}\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
