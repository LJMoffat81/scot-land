"""HOLDOUT-UR6-2 trailing n25i rows, council match, and the no-price stop.

SYNTHETIC labelled rows. No 2024 price cell is read.
"""

from __future__ import annotations

import ast
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from georgist.v14_real_headers import REAL_HEADERS
from scripts import build_v14_drc_components as components
from scripts import run_v14_no_price as no_price_driver
from scripts.build_v14_check_sources import PIN_TABLE as SOURCE_PINS
from scripts.build_v14_check_sources import main as source_main
from scripts.build_v14_drc_components import (
    PIN_TABLE,
    build_components,
    copy_pin_table,
    load_n25i,
    sha256_file,
    write_components,
)
from scripts.build_v14_drc_components import main as components_main

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "v14_pr3a"
N25I_PIN = "6fdf2951223dd01e5790069d61428ce25d92aed5f260197ee546223113df4e49"
WIDTH = len(REAL_HEADERS["n25i"])


def _paths() -> dict[str, Path]:
    return {
        "epc": FIXTURES / "epc.csv",
        "nrs2017": FIXTURES / "nrs2017.csv",
        "plot_bands": FIXTURES / "plot_bands.csv",
        "ur6": FIXTURES / "ur6.csv",
        "n25i": FIXTURES / "n25i.xlsx",
    }


def _pins(paths: dict[str, Path]) -> dict[str, str]:
    pins = copy_pin_table()
    for name, path in paths.items():
        if name in pins and path.is_file():
            pins[name] = sha256_file(path)
    return pins


def _zone(index: int) -> str:
    return f"S01{index:06d}"


def _council(index: int) -> str:
    return f"S12{(index % 32) + 1:06d}"


def _cells(code: object, council: object) -> list[object]:
    row: list[object] = []
    for column in REAL_HEADERS["n25i"]:
        if column == "Data Zone code":
            row.append(code)
        elif column == "Council area code":
            row.append(council)
        elif column in {"Data zone name", "Council area name"}:
            row.append("SYNTHETIC")
        elif column == "Total number of dwellings":
            row.append("10")
        else:
            row.append("0")
    return row


def _save(path: Path, sheet) -> None:
    sheet.title = "2025"
    path.parent.mkdir(parents=True, exist_ok=True)
    sheet.parent.save(path)


def _workbook(header_row: int = 1) -> tuple[Workbook, object]:
    workbook = Workbook()
    sheet = workbook.active
    if header_row != 1:
        raise AssertionError(header_row)
    sheet.append(["SYNTHETIC"])
    for _ in range(3):
        sheet.append([None])
    return workbook, sheet


def _write_rows(path: Path, rows: list[list[object]], *, preamble: bool = True) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "2025"
    if preamble:
        sheet.append(["SYNTHETIC"])
        sheet.append([None])
        sheet.append([None])
        sheet.append([None])
    sheet.append(list(REAL_HEADERS["n25i"]))
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _pair_row(index: int) -> list[object]:
    return _cells(_zone(index), _council(index))


def _csv(path: Path, header: list[str], rows: list[list[str]]) -> None:
    lines = ["# SYNTHETIC", ",".join(header)]
    lines.extend(",".join(cells) for cells in rows)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _row(name: str, values: dict[str, str]) -> list[str]:
    return [values.get(column, "") for column in REAL_HEADERS[name]]


def _component_paths(tmp_path: Path, n25i: Path) -> dict[str, Path]:
    paths = dict(_paths())
    paths["n25i"] = n25i
    for key in ("epc", "nrs2017", "plot_bands", "ur6"):
        path = tmp_path / f"{key}.csv"
        path.write_bytes(paths[key].read_bytes())
        paths[key] = path
    paths["n25i"] = n25i
    return paths


def test_u1_trailing_empty_row_loads(tmp_path: Path) -> None:
    """SYNTHETIC. 6976 data rows and one trailing empty row load."""
    path = tmp_path / "n25i.xlsx"
    rows = [_pair_row(index) for index in range(1, 6977)]
    rows.append([""] * WIDTH)
    _write_rows(path, rows, preamble=False)
    loaded = load_n25i(path, expected_rows=6976)
    assert len(loaded) == 6976
    assert loaded[_zone(6976)] == _council(6976)


def test_u2_three_trailing_none_rows(tmp_path: Path) -> None:
    """SYNTHETIC. Three trailing all-None rows are not counted."""
    path = tmp_path / "n25i.xlsx"
    rows = [_pair_row(index) for index in range(1, 6977)]
    rows.extend([""] * WIDTH for _ in range(3))
    _write_rows(path, rows, preamble=False)
    loaded = load_n25i(path, expected_rows=6976)
    assert len(loaded) == 6976


def test_u2b_trailing_whitespace_refuses(tmp_path: Path) -> None:
    """SYNTHETIC. A trailing whitespace row is a blank Data Zone code."""
    path = tmp_path / "n25i.xlsx"
    rows = [_pair_row(index) for index in range(1, 6977)]
    rows.append([" "] * WIDTH)
    _write_rows(path, rows, preamble=False)
    with pytest.raises(Exception) as raised:
        load_n25i(path, expected_rows=6976)
    assert str(raised.value) == (
        "B1_COMPONENTS: ERROR: Data Zone code row 6978: expected value, got blank"
    )


def test_u3_empty_before_data(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """SYNTHETIC. An all-None row before a later data row is refused."""
    calls = {"down": 0}
    monkeypatch.setattr(
        "scripts.build_v14_drc_components.borrow_plot_bands",
        lambda *_a, **_k: calls.__setitem__("down", calls["down"] + 1),
    )
    path = tmp_path / "n25i.xlsx"
    rows = [_pair_row(1), _pair_row(2), _pair_row(3), [""] * WIDTH, _pair_row(4)]
    _write_rows(path, rows, preamble=False)
    paths = _component_paths(tmp_path, path)
    out = tmp_path / "components.json"
    with pytest.raises(Exception) as raised:
        build_components(paths, pins=_pins(paths))
    assert str(raised.value) == (
        "B1_COMPONENTS: ERROR: n25i sheet 2025 row 5: empty row before data row 6"
    )
    assert calls["down"] == 0
    assert not out.exists()


def test_u3b_whitespace_gap_is_blank_code(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """SYNTHETIC. Whitespace is not an empty row, so the per-row check fires."""
    calls = {"down": 0}
    monkeypatch.setattr(
        "scripts.build_v14_drc_components.borrow_plot_bands",
        lambda *_a, **_k: calls.__setitem__("down", calls["down"] + 1),
    )
    path = tmp_path / "n25i.xlsx"
    rows = [_pair_row(1), _pair_row(2), _pair_row(3), [" "] * WIDTH, _pair_row(4)]
    _write_rows(path, rows, preamble=False)
    paths = _component_paths(tmp_path, path)
    out = tmp_path / "components.json"
    with pytest.raises(Exception) as raised:
        build_components(paths, pins=_pins(paths))
    assert str(raised.value) == (
        "B1_COMPONENTS: ERROR: Data Zone code row 5: expected value, got blank"
    )
    assert calls["down"] == 0
    assert not out.exists()


def test_u4_blank_code_middle(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """SYNTHETIC. A blank code with a council present keeps the blank-code refusal."""
    calls = {"down": 0}
    monkeypatch.setattr(
        "scripts.build_v14_drc_components.borrow_plot_bands",
        lambda *_a, **_k: calls.__setitem__("down", calls["down"] + 1),
    )
    path = tmp_path / "n25i.xlsx"
    rows = [_pair_row(1), _cells("", "S12000033")]
    _write_rows(path, rows, preamble=False)
    paths = _component_paths(tmp_path, path)
    out = tmp_path / "components.json"
    with pytest.raises(Exception) as raised:
        build_components(paths, pins=_pins(paths))
    assert str(raised.value) == (
        "B1_COMPONENTS: ERROR: Data Zone code row 3: expected value, got blank"
    )
    assert calls["down"] == 0
    assert not out.exists()


def test_u4b_last_blank_not_dropped(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """SYNTHETIC. A blank last data row is not treated as a trailing empty row."""
    calls = {"down": 0}
    monkeypatch.setattr(
        "scripts.build_v14_drc_components.borrow_plot_bands",
        lambda *_a, **_k: calls.__setitem__("down", calls["down"] + 1),
    )
    path = tmp_path / "n25i.xlsx"
    rows = [_pair_row(1), _pair_row(2), _cells("", "S12000033"), [""] * WIDTH, [""] * WIDTH]
    _write_rows(path, rows, preamble=False)
    paths = _component_paths(tmp_path, path)
    out = tmp_path / "components.json"
    with pytest.raises(Exception) as raised:
        build_components(paths, pins=_pins(paths))
    assert str(raised.value) == (
        "B1_COMPONENTS: ERROR: Data Zone code row 4: expected value, got blank"
    )
    assert calls["down"] == 0
    assert not out.exists()


def test_u4c_whitespace_code(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """SYNTHETIC. A whitespace code is blank at that sheet row."""
    calls = {"down": 0}
    monkeypatch.setattr(
        "scripts.build_v14_drc_components.borrow_plot_bands",
        lambda *_a, **_k: calls.__setitem__("down", calls["down"] + 1),
    )
    path = tmp_path / "n25i.xlsx"
    rows = [_pair_row(1), _cells("  ", "S12000033")]
    _write_rows(path, rows, preamble=False)
    paths = _component_paths(tmp_path, path)
    out = tmp_path / "components.json"
    with pytest.raises(Exception) as raised:
        build_components(paths, pins=_pins(paths))
    assert str(raised.value) == (
        "B1_COMPONENTS: ERROR: Data Zone code row 3: expected value, got blank"
    )
    assert calls["down"] == 0
    assert not out.exists()


def _pin_sha(monkeypatch: pytest.MonkeyPatch) -> None:
    real = sha256_file

    def fake(candidate: Path) -> str:
        candidate = Path(candidate)
        if candidate.suffix == ".xlsx":
            return str(PIN_TABLE["n25i"])
        if candidate.stem in PIN_TABLE:
            return str(PIN_TABLE[candidate.stem])
        return real(candidate)

    monkeypatch.setattr("scripts.build_v14_drc_components.sha256_file", fake)


def test_u5_pin_count_6975_and_6977(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """SYNTHETIC. A production-pin sha applies the 6976 row count."""
    calls = {"down": 0}
    monkeypatch.setattr(
        "scripts.build_v14_drc_components.borrow_plot_bands",
        lambda *_a, **_k: calls.__setitem__("down", calls["down"] + 1),
    )

    def run(count: int) -> str:
        path = tmp_path / f"n25i-{count}.xlsx"
        rows = [_pair_row(index) for index in range(1, count + 1)]
        if count == 6975:
            rows.append([""] * WIDTH)
        _write_rows(path, rows, preamble=False)
        out = tmp_path / f"out-{count}.json"
        code = components_main(
            [
                "--out",
                str(out),
                "--epc",
                str(FIXTURES / "epc.csv"),
                "--nrs2017",
                str(FIXTURES / "nrs2017.csv"),
                "--plot-bands",
                str(FIXTURES / "plot_bands.csv"),
                "--ur6",
                str(FIXTURES / "ur6.csv"),
                "--n25i",
                str(path),
            ]
        )
        assert code == 1
        assert not out.exists()
        return path.name

    _pin_sha(monkeypatch)
    captured = []

    real_write = __import__("sys").stderr.write

    def grab(text: str) -> int:
        captured.append(text)
        return real_write(text)

    monkeypatch.setattr("sys.stderr.write", grab)
    run(6975)
    assert captured[-1].strip() == "B1_COMPONENTS: ERROR: n25i rows: expected 6976, got 6975"
    captured.clear()
    run(6977)
    assert captured[-1].strip() == "B1_COMPONENTS: ERROR: n25i rows: expected 6976, got 6977"
    assert calls["down"] == 0


def test_u5a_cli_expected_rows(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """SYNTHETIC. The CLI passes 6976 only when the file sha is the production pin."""
    assert components.N25I_EXPECTED_ROWS == 6976
    assert components.PIN_TABLE["n25i"] == N25I_PIN
    seen: list[int | None] = []

    def spy(path: Path, *, expected_rows: int | None = None) -> dict[str, str]:
        seen.append(expected_rows)
        return {_zone(1): _council(1)}

    monkeypatch.setattr(components, "load_n25i", spy)
    monkeypatch.setattr(components, "borrow_plot_bands", lambda *_a, **_k: {})
    real = components.sha256_file

    def matching(candidate: Path) -> str:
        stem = Path(candidate).stem
        if stem in components.PIN_TABLE:
            return str(components.PIN_TABLE[stem])
        return real(Path(candidate))

    monkeypatch.setattr(components, "sha256_file", matching)
    argv = [
        "--out",
        str(tmp_path / "out.json"),
        "--epc",
        str(FIXTURES / "epc.csv"),
        "--nrs2017",
        str(FIXTURES / "nrs2017.csv"),
        "--plot-bands",
        str(FIXTURES / "plot_bands.csv"),
        "--ur6",
        str(FIXTURES / "ur6.csv"),
        "--n25i",
        str(FIXTURES / "n25i.xlsx"),
    ]
    components.main(argv)
    assert seen == [6976]
    seen.clear()

    def other(candidate: Path) -> str:
        stem = Path(candidate).stem
        if stem == "n25i":
            return "ab" * 32
        if stem in components.PIN_TABLE:
            return str(components.PIN_TABLE[stem])
        return real(Path(candidate))

    monkeypatch.setattr(components, "sha256_file", other)
    monkeypatch.setattr(components, "assert_pins", lambda *_a, **_k: None)
    components.main(argv)
    assert seen == [None]


def test_u5b_injected_count(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """SYNTHETIC. An injected count accepts 3 rows and refuses 2."""
    calls = {"down": 0}
    monkeypatch.setattr(
        "scripts.build_v14_drc_components.borrow_plot_bands",
        lambda *_a, **_k: calls.__setitem__("down", calls["down"] + 1),
    )
    three = tmp_path / "three.xlsx"
    _write_rows(three, [_pair_row(index) for index in range(1, 4)], preamble=False)
    loaded = load_n25i(three, expected_rows=3)
    assert list(loaded) == [_zone(1), _zone(2), _zone(3)]
    two = tmp_path / "two.xlsx"
    _write_rows(two, [_pair_row(1), _pair_row(2)], preamble=False)
    out = tmp_path / "out.json"
    with pytest.raises(Exception) as raised:
        load_n25i(two, expected_rows=3)
    assert str(raised.value) == "B1_COMPONENTS: ERROR: n25i rows: expected 3, got 2"
    assert calls["down"] == 0
    assert not out.exists()


def test_u5c_none_skips_count(tmp_path: Path) -> None:
    """SYNTHETIC. expected_rows None does not apply the production count."""
    path = tmp_path / "n25i.xlsx"
    _write_rows(path, [_pair_row(1), _pair_row(2)], preamble=False)
    assert sha256_file(path) != N25I_PIN
    loaded = load_n25i(path, expected_rows=None)
    assert len(loaded) == 2


class _Frame:
    columns = ["code", "council_code"]

    def __init__(self, pairs: list[tuple[str, str]]) -> None:
        self._pairs = pairs

    def __getitem__(self, key: str) -> list[str]:
        if key == "code":
            return [code for code, _council in self._pairs]
        if key == "council_code":
            return [council for _code, council in self._pairs]
        raise KeyError(key)


def _mismatch(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    validated: list[tuple[str, str]],
    harvested: list[tuple[str, str]],
) -> str:
    calls = {"down": 0}
    monkeypatch.setattr(
        "scripts.build_v14_drc_components.borrow_plot_bands",
        lambda *_a, **_k: calls.__setitem__("down", calls["down"] + 1),
    )
    monkeypatch.setattr(
        "scripts.build_v14_drc_components.hv.load_nrs_dwellings",
        lambda *_a, **_k: _Frame(harvested),
    )
    path = tmp_path / "n25i.xlsx"
    _write_rows(path, [_cells(code, council) for code, council in validated], preamble=True)
    paths = _component_paths(tmp_path, path)
    out = tmp_path / "components.json"
    with pytest.raises(Exception) as raised:
        records = build_components(paths, pins=_pins(paths))
        write_components(out, records)
    assert calls["down"] == 0
    assert not out.exists()
    return str(raised.value)


def test_u6_one_changed_council(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """SYNTHETIC. One harvester council difference is refused."""
    message = _mismatch(
        tmp_path,
        monkeypatch,
        [(_zone(1), "S12000033")],
        [(_zone(1), "S12000034")],
    )
    assert message == (
        "B1_COMPONENTS: ERROR: n25i council mismatch at S01000001: "
        "validated 'S12000033', harvester 'S12000034'"
    )


def test_u6_equal_output_bytes(tmp_path: Path) -> None:
    """SYNTHETIC. Equal council maps keep the pre-change component bytes."""
    spec = [
        ("S01000001", "S12000001", "1"),
        ("S01000002", "S12000002", "2"),
        ("S01000003", "S12000001", "1"),
    ]
    paths = {}
    for key in ("epc", "nrs2017", "plot_bands", "ur6"):
        path = tmp_path / f"{key}.csv"
        header = list(REAL_HEADERS[key])
        if key == "ur6":
            rows = []
            for code, _council, ur6_class in spec:
                values = {column: "" for column in header}
                values["code"] = code
                values["ur6_class"] = ur6_class
                values["ur6_name"] = "SYNTHETIC"
                values["source"] = "SYNTHETIC"
                values["licence"] = "OGL-3.0"
                rows.append(_row(key, values))
        elif key == "epc":
            rows = []
            for code, _council, _ur6_class in spec:
                values = {column: "" for column in header}
                values["code"] = code
                for column in header:
                    if column.startswith("age_") or column.startswith("recent10_age_"):
                        values[column] = "1" if column == "age_pre1919_share" else "0"
                rows.append(_row(key, values))
        elif key == "nrs2017":
            rows = []
            for code, _council, _ur6_class in spec:
                values = {column: "" for column in header}
                values["code"] = code
                for column in header:
                    if column.startswith("type_"):
                        values[column] = "10" if column == "type_detached" else "0"
                rows.append(_row(key, values))
        else:
            rows = []
            bands = [
                {"all_ages_share_lt005": "0.5", "all_ages_share_005_010": "0.25", "all_ages_share_gt010": "0.25", "n_detached_all_ages_matched": "4"},
                {"all_ages_share_lt005": "0", "all_ages_share_005_010": "0.5", "all_ages_share_gt010": "0.5", "n_detached_all_ages_matched": "2"},
                {},
            ]
            for (code, _council, _ur6_class), band in zip(spec, bands):
                values = {column: "" for column in header}
                values["code"] = code
                values.update(band)
                rows.append(_row(key, values))
        _csv(path, header, rows)
        paths[key] = path
    n25i = tmp_path / "n25i.xlsx"
    _write_rows(n25i, [_cells(code, council) for code, council, _ur6 in spec], preamble=True)
    paths["n25i"] = n25i
    records = build_components(paths, pins=_pins(paths))
    out = tmp_path / "components.json"
    write_components(out, records)
    expected = [
        {
            "dz2011": "S01000001",
            "council_code": "S12000001",
            "epc_row": {column: records[0]["epc_row"][column] for column in records[0]["epc_row"]},
            "nrs2017_row": records[0]["nrs2017_row"],
            "estate_row": {**records[0]["estate_row"]},
        },
        {
            "dz2011": "S01000002",
            "council_code": "S12000002",
            "epc_row": records[1]["epc_row"],
            "nrs2017_row": records[1]["nrs2017_row"],
            "estate_row": records[1]["estate_row"],
        },
        {
            "dz2011": "S01000003",
            "council_code": "S12000001",
            "epc_row": records[2]["epc_row"],
            "nrs2017_row": records[2]["nrs2017_row"],
            "estate_row": {
                "code": "S01000003",
                "all_ages_share_lt005": 0.5,
                "all_ages_share_005_010": 0.25,
                "all_ages_share_gt010": 0.25,
                "_plot_bands_source": "borrowed",
                "_borrow_cell": "('S12000001', 1)",
                "_borrow_n_zones": 1,
            },
        },
    ]
    assert out.read_bytes() == (json.dumps(expected, indent=2, sort_keys=True) + "\n").encode("utf-8")


def test_u6b_case_difference(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """SYNTHETIC. Council comparison does not casefold."""
    message = _mismatch(
        tmp_path,
        monkeypatch,
        [(_zone(1), "S12000033")],
        [(_zone(1), "s12000033")],
    )
    assert message == (
        "B1_COMPONENTS: ERROR: n25i council mismatch at S01000001: "
        "validated 'S12000033', harvester 's12000033'"
    )


def test_u6c_trailing_space(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """SYNTHETIC. A trailing space in the harvester council is kept."""
    message = _mismatch(
        tmp_path,
        monkeypatch,
        [(_zone(1), "S12000033")],
        [(_zone(1), "S12000033 ")],
    )
    assert message == (
        "B1_COMPONENTS: ERROR: n25i council mismatch at S01000001: "
        "validated 'S12000033', harvester 'S12000033 '"
    )


def test_u6d_harvester_only(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """SYNTHETIC. A zone only in the harvester map is absent on the validated side."""
    message = _mismatch(
        tmp_path,
        monkeypatch,
        [(_zone(1), "S12000033")],
        [(_zone(1), "S12000033"), (_zone(2), "S12000033")],
    )
    assert message == (
        "B1_COMPONENTS: ERROR: n25i council mismatch at S01000002: "
        "validated absent, harvester 'S12000033'"
    )


def test_u6e_validated_only(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """SYNTHETIC. A zone only in the validated map is absent on the harvester side."""
    message = _mismatch(
        tmp_path,
        monkeypatch,
        [(_zone(1), "S12000033")],
        [],
    )
    assert message == (
        "B1_COMPONENTS: ERROR: n25i council mismatch at S01000001: "
        "validated 'S12000033', harvester absent"
    )


def test_u6f_first_sorted_mismatch(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """SYNTHETIC. Two mismatches are reported at the first sorted zone."""
    message = _mismatch(
        tmp_path,
        monkeypatch,
        [(_zone(2), "S12000033"), (_zone(1), "S12000033")],
        [(_zone(2), "S12000034"), (_zone(1), "S12000034")],
    )
    assert message == (
        "B1_COMPONENTS: ERROR: n25i council mismatch at S01000001: "
        "validated 'S12000033', harvester 'S12000034'"
    )


def test_u7_no_harvester_assignment() -> None:
    """SYNTHETIC. load_n25i does not assign a harvester column into councils."""
    source = Path(__file__).resolve().parents[1] / "scripts" / "build_v14_drc_components.py"
    tree = ast.parse(source.read_text(encoding="utf-8"))
    function = next(node for node in tree.body if isinstance(node, ast.FunctionDef) and node.name == "load_n25i")
    for node in ast.walk(function):
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute):
            assert node.func.attr != "load_nrs_dwellings"
        if isinstance(node, ast.Attribute):
            assert node.attr != "council_code"
        if isinstance(node, ast.Assign):
            for target in node.targets:
                assert not isinstance(target, ast.Subscript)


def _rent_file(path: Path, pending: int, not_applicable: int) -> None:
    header = list(REAL_HEADERS["rent_table"])
    rows = []
    observed = {column: "1" for column in header}
    observed["zone_id"] = "Z1"
    observed["status"] = "observed_pooled"
    observed["central"] = "4"
    observed["low"] = "1"
    rows.append([observed[column] for column in header])
    for index in range(pending):
        item = {column: "1" for column in header}
        item["zone_id"] = f"P{index}"
        item["status"] = "model_fill_pending"
        rows.append([item[column] for column in header])
    for index in range(not_applicable):
        item = {column: "1" for column in header}
        item["zone_id"] = f"N{index}"
        item["status"] = "not_applicable"
        rows.append([item[column] for column in header])
    _csv(path, header, rows)


def test_u8_check3_not_run_year_gate(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]) -> None:
    """SYNTHETIC. Check 3 NOT_RUN is non-gating in 2023 and gating in 2024."""
    plots = tmp_path / "plots.csv"
    _csv(plots, list(REAL_HEADERS["plot_bands"]), [["SYNTHETIC"] * len(REAL_HEADERS["plot_bands"])])
    houses = tmp_path / "houses.csv"
    _csv(houses, ["code", "house_plot_m2"], [["Z1", "1"]])
    rent = tmp_path / "rent.csv"
    _rent_file(rent, 129, 3)
    monkeypatch.setitem(SOURCE_PINS, "plot_bands", sha256_file(plots))
    monkeypatch.setitem(SOURCE_PINS, "gis_house_plot", sha256_file(houses))
    monkeypatch.setitem(SOURCE_PINS, "rent_table", sha256_file(rent))
    check3 = tmp_path / "check3.json"
    check4 = tmp_path / "check4.json"
    argv = [
        "--year",
        "2023",
        "--no-price",
        "--plot-bands",
        str(plots),
        "--gis-house-plot",
        str(houses),
        "--check3-out",
        str(check3),
        "--rent-table",
        str(rent),
        "--check4-out",
        str(check4),
    ]
    assert source_main(argv) == 0
    out = capsys.readouterr().out
    assert (
        "B5_CHECK4: excluded=132\n"
        "CHECK4_EXCLUDED: total=132 model_fill_pending=129 not_applicable=3\n"
    ) in out
    assert "B5_CHECK3: NOT_RUN: plot_bands header: expected pass_20, got absent" in out
    bad = tmp_path / "bad_rent.csv"
    _rent_file(bad, 1, 1)
    monkeypatch.setitem(SOURCE_PINS, "rent_table", sha256_file(bad))
    failed = argv.copy()
    failed[failed.index(str(rent))] = str(bad)
    assert source_main(["--year", "2023", "--no-price", "--rent-table", str(bad), "--check4-out", str(tmp_path / "c4.json")]) == 1
    assert source_main(["--year", "2024", "--no-price", "--plot-bands", str(plots), "--gis-house-plot", str(houses), "--check3-out", str(tmp_path / "c3.json")]) == 1


def _no_price_root(root: Path) -> None:
    code = "S01000001"
    council = "S12000001"
    for key, filler in (
        ("epc", lambda column, values: values.update({column: "1" if column == "age_pre1919_share" else "0"}) if column.startswith("age_") or column.startswith("recent10_age_") else None),
        ("nrs2017", lambda column, values: values.update({column: "10" if column == "type_detached" else "0"}) if column.startswith("type_") else None),
        ("plot_bands", lambda column, values: values.update({"all_ages_share_lt005": "0.5", "all_ages_share_005_010": "0.25", "all_ages_share_gt010": "0.25", "n_detached_all_ages_matched": "4"})),
        ("ur6", lambda column, values: values.update({"ur6_class": "1", "ur6_name": "SYNTHETIC", "source": "SYNTHETIC", "licence": "OGL-3.0"})),
    ):
        header = list(REAL_HEADERS[key])
        values = {column: "" for column in header}
        values["code"] = code
        for column in header:
            filler(column, values)
        _csv(root / f"{key}.csv", header, [_row(key, values)])
    _write_rows(root / "n25i.xlsx", [_cells(code, council)], preamble=True)
    _csv(root / "check3_plots.csv", list(REAL_HEADERS["plot_bands"]), [["SYNTHETIC"] * len(REAL_HEADERS["plot_bands"])])
    _csv(root / "gis_house_plot.csv", ["code", "house_plot_m2"], [[code, "1"]])
    _rent_file(root / "rent_table.csv", 129, 3)
    (root / "house_sales.csv").write_text("# SYNTHETIC\n", encoding="utf-8")


def test_u9_no_price_chain_stops(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """SYNTHETIC. The no-price chain stops after check_sources and opens no house sales."""
    driver = no_price_driver
    calls = {"check_inputs": 0, "fit": 0, "primary_2023": 0, "score": 0, "check_sources": 0}

    def _count(name: str):
        def _inner(*_args, **_kwargs):
            calls[name] += 1
            return 0

        return _inner

    monkeypatch.setattr(driver, "check_inputs_main", _count("check_inputs"))
    monkeypatch.setattr(driver, "fit_main", _count("fit"))
    monkeypatch.setattr(driver, "primary_2023_main", _count("primary_2023"))
    monkeypatch.setattr(driver, "score_main", _count("score"))
    real_sources = driver.check_sources_main

    def _sources(argv=None):
        calls["check_sources"] += 1
        return real_sources(argv)

    monkeypatch.setattr(driver, "check_sources_main", _sources)
    root = tmp_path / "run"
    root.mkdir()
    _no_price_root(root)
    for name in ("epc", "nrs2017", "plot_bands", "ur6", "n25i"):
        monkeypatch.setitem(PIN_TABLE, name, sha256_file(root / ("n25i.xlsx" if name == "n25i" else f"{name}.csv")))
    monkeypatch.setitem(SOURCE_PINS, "plot_bands", sha256_file(root / "check3_plots.csv"))
    monkeypatch.setitem(SOURCE_PINS, "gis_house_plot", sha256_file(root / "gis_house_plot.csv"))
    monkeypatch.setitem(SOURCE_PINS, "rent_table", sha256_file(root / "rent_table.csv"))
    opened: list[str] = []
    real_open = Path.open

    def tracking(self: Path, *args, **kwargs):
        if "house_sales" in self.name:
            opened.append(self.name)
        return real_open(self, *args, **kwargs)

    monkeypatch.setattr(Path, "open", tracking)
    code = driver.main(["--root", str(root), "--year", "2023", "--no-price"])
    assert code == 0
    assert calls == {"check_inputs": 0, "fit": 0, "primary_2023": 0, "score": 0, "check_sources": 1}
    assert opened == []


def test_u12_pr3a_synthetic_marker() -> None:
    """SYNTHETIC. The pr3a n25i fixture is labelled on sheet 2025 cell A1."""
    path = FIXTURES / "n25i.xlsx"
    workbook = load_workbook(path, data_only=True)
    try:
        assert workbook.sheetnames == ["2025"]
        assert workbook["2025"]["A1"].value == "SYNTHETIC"
    finally:
        workbook.close()
